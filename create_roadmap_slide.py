"""
Generates a detailed Problem Statement & Project Evolution roadmap slide.
Run with:  venv\Scripts\python.exe create_roadmap_slide.py
Output:    Roadmap_Evolution_Slide.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import pptx.oxml.ns as nsmap
from lxml import etree

# ── Colour palette ─────────────────────────────────────────────────────────────
C_BG_DARK    = RGBColor(0x0D, 0x11, 0x17)
C_BG_CARD    = RGBColor(0x16, 0x21, 0x3E)
C_ACCENT     = RGBColor(0x63, 0xB3, 0xED)   # sky blue
C_ACCENT2    = RGBColor(0x76, 0xE4, 0xF7)   # teal
C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT_TEXT = RGBColor(0xA8, 0xB2, 0xD8)
C_MUTED      = RGBColor(0x71, 0x80, 0x96)
C_GREEN      = RGBColor(0x68, 0xD3, 0x91)
C_PURPLE     = RGBColor(0xB7, 0x94, 0xF4)
C_ORANGE     = RGBColor(0xF6, 0xAD, 0x55)
C_RED        = RGBColor(0xFC, 0x6B, 0x6B)
C_YELLOW     = RGBColor(0xF6, 0xE0, 0x5A)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)


def new_prs():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def fill_bg(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height,
             fill_color=None, line_color=None, line_width=Pt(0)):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.line.width = line_width
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_text_box(slide, text, left, top, width, height,
                 font_size=Pt(14), bold=False, color=C_WHITE,
                 align=PP_ALIGN.LEFT, wrap=True, italic=False):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txBox


def add_multi_line_text(slide, lines, left, top, width, height,
                        font_size=Pt(10.5), color=C_LIGHT_TEXT,
                        bold=False, line_spacing_pt=5):
    """Add a text box with multiple paragraphs (one per line in list)."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(line_spacing_pt)
        run = p.add_run()
        run.text = line
        run.font.size = font_size
        run.font.bold = bold
        run.font.color.rgb = color
    return txBox


def add_gradient_rect(slide, left, top, width, height,
                      color1: RGBColor, color2: RGBColor, angle=5400000):
    shape = add_rect(slide, left, top, width, height, fill_color=color1)
    sp = shape._element
    spPr = sp.find(nsmap.qn('p:spPr'))
    if spPr is None:
        return shape
    solid_fill = spPr.find(nsmap.qn('a:solidFill'))
    if solid_fill is not None:
        spPr.remove(solid_fill)
    grad_xml = f"""
    <a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1">
      <a:gsLst>
        <a:gs pos="0"><a:srgbClr val="{str(color1)}"/></a:gs>
        <a:gs pos="100000"><a:srgbClr val="{str(color2)}"/></a:gs>
      </a:gsLst>
      <a:lin ang="{angle}" scaled="0"/>
    </a:gradFill>"""
    spPr.append(etree.fromstring(grad_xml))
    return shape


# ─────────────────────────────────────────────────────────────────────────────
#  SLIDE: Problem Statement & Project Evolution
# ─────────────────────────────────────────────────────────────────────────────

def build_roadmap_slide(prs):
    slide = blank_slide(prs)
    fill_bg(slide, C_BG_DARK)

    # ── Full background gradient ──────────────────────────────────────────────
    add_gradient_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H,
                      RGBColor(0x08, 0x0C, 0x16), RGBColor(0x09, 0x1E, 0x3A))

    # Decorative top-right glow
    add_gradient_rect(slide, Inches(10.0), Inches(0), Inches(3.33), Inches(2.5),
                      RGBColor(0x10, 0x2A, 0x50), RGBColor(0x08, 0x0C, 0x16))

    # ── Top accent bar ────────────────────────────────────────────────────────
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), fill_color=C_ACCENT)

    # ── Header band ──────────────────────────────────────────────────────────
    add_rect(slide, Inches(0), Inches(0.06), SLIDE_W, Inches(0.88),
             fill_color=RGBColor(0x0C, 0x18, 0x2E))

    # Pill icon
    add_rect(slide, Inches(0.28), Inches(0.13), Inches(0.65), Inches(0.65),
             fill_color=C_ACCENT)
    add_text_box(slide, "🗺️", Inches(0.28), Inches(0.11), Inches(0.65), Inches(0.70),
                 font_size=Pt(24), align=PP_ALIGN.CENTER)

    # Main title
    add_text_box(slide, "Problem Statement & Project Evolution Roadmap",
                 Inches(1.08), Inches(0.10), Inches(9.5), Inches(0.48),
                 font_size=Pt(23), bold=True, color=C_WHITE)
    add_text_box(slide,
                 "From synthetic look-alike data  →  PII-grounded generation  →  Pattern-driven production-ready data",
                 Inches(1.08), Inches(0.57), Inches(9.8), Inches(0.32),
                 font_size=Pt(10.5), color=C_ACCENT, italic=True)

    # Right badge
    add_rect(slide, Inches(11.0), Inches(0.20), Inches(2.05), Inches(0.52),
             fill_color=RGBColor(0x14, 0x22, 0x40), line_color=C_ACCENT, line_width=Pt(1))
    add_text_box(slide, "⚡  Databricks · Claude Sonnet",
                 Inches(11.0), Inches(0.20), Inches(2.05), Inches(0.52),
                 font_size=Pt(8.5), bold=True, color=C_ACCENT2, align=PP_ALIGN.CENTER)

    # ── Thin separator ────────────────────────────────────────────────────────
    add_rect(slide, Inches(0), Inches(0.94), SLIDE_W, Inches(0.022),
             fill_color=RGBColor(0x1C, 0x38, 0x60))

    # ══════════════════════════════════════════════════════════════════════════
    #  PROBLEM STATEMENT CARD  (full-width, compact)
    # ══════════════════════════════════════════════════════════════════════════
    PS_TOP  = Inches(1.00)
    PS_H    = Inches(1.22)
    PS_L    = Inches(0.22)
    PS_W    = SLIDE_W - Inches(0.44)

    add_rect(slide, PS_L, PS_TOP, PS_W, PS_H,
             fill_color=RGBColor(0x12, 0x1E, 0x38),
             line_color=C_RED, line_width=Pt(1.4))
    # Red left accent bar
    add_rect(slide, PS_L, PS_TOP, Inches(0.045), PS_H, fill_color=C_RED)
    # Top colour wash
    add_rect(slide, PS_L, PS_TOP, PS_W, Inches(0.04), fill_color=C_RED)

    add_text_box(slide, "⚠️  The Problem  —  Why We Started This Initiative",
                 PS_L + Inches(0.13), PS_TOP + Inches(0.06),
                 PS_W - Inches(0.2), Inches(0.34),
                 font_size=Pt(12.5), bold=True, color=C_RED)

    problem_points = [
        "🔐  Production data contains PII — sharing it across dev, QA, and vendor teams violates GDPR, HIPAA, and internal data-governance policies.",
        "⏳  Access approvals take weeks — teams are blocked waiting for sanitised extracts before pipelines, models, or dashboards can be built.",
        "📉  Real data lacks edge cases — rare events, outliers, and stress-test scenarios needed for robust testing are absent from standard exports.",
        "💸  Manual data crafting is slow & error-prone — hand-building realistic datasets doesn't scale and introduces subtle inconsistencies.",
    ]

    # Lay out problem points in two columns
    for idx, pt in enumerate(problem_points):
        col = idx % 2
        row = idx // 2
        lx = PS_L + Inches(0.15) + col * Inches(6.55)
        ty = PS_TOP + Inches(0.44) + row * Inches(0.36)
        add_text_box(slide, pt, lx, ty, Inches(6.3), Inches(0.34),
                     font_size=Pt(9.5), color=C_LIGHT_TEXT)

    # ══════════════════════════════════════════════════════════════════════════
    #  THREE PHASE TILES
    # ══════════════════════════════════════════════════════════════════════════
    PHASE_TOP  = PS_TOP + PS_H + Inches(0.12)
    PHASE_H    = SLIDE_H - PHASE_TOP - Inches(0.28)
    TILE_W     = (SLIDE_W - Inches(0.55)) / 3
    GAP        = Inches(0.08)

    phase_colors = [C_GREEN, C_ACCENT, C_PURPLE]
    phase_bg     = [RGBColor(0x0D, 0x22, 0x1A),
                    RGBColor(0x0C, 0x1C, 0x38),
                    RGBColor(0x1A, 0x0D, 0x30)]

    # ── Connector arrows between tiles ────────────────────────────────────────
    for i in range(2):
        arrow_l = Inches(0.22) + (i + 1) * TILE_W + i * GAP
        arrow_t = PHASE_TOP + PHASE_H / 2 - Inches(0.2)
        add_text_box(slide, "▶", arrow_l, arrow_t,
                     GAP, Inches(0.4), font_size=Pt(14),
                     color=C_MUTED, align=PP_ALIGN.CENTER)

    phases = [
        # ── PHASE 1 ──────────────────────────────────────────────────────────
        {
            "num":    "Phase 1",
            "title":  "Look-alike Data Generation",
            "status": "✅  COMPLETED",
            "status_color": C_GREEN,
            "icon":   "🧪",
            "color":  C_GREEN,
            "bg":     phase_bg[0],
            "what": "Build a fully automated LLM-driven agent that generates "
                    "synthetic data resembling real datasets — with no dependency "
                    "on actual production tables.",
            "bullets": [
                ("🤖", "LLM Prompting",
                 "Designed structured prompt templates for Claude Sonnet (Databricks) "
                 "to produce schema-compliant rows with realistic field distributions."),
                ("📦", "Cluster (CSV) Schema",
                 "STORE_NO · DEPT_NO · CSTD_GRADE · PHASE_STARTDATE — "
                 "unique combos per batch, zero-padded formats enforced."),
                ("🛍️", "ProductVendor (JSON) Schema",
                 "Nested productId · season · vendors[] objects with 18-digit IDs, "
                 "supplier codes (M-prefix), factory numbers & isActive flags."),
                ("🔄", "Batch Streaming + Retry",
                 "Large volumes split into configurable batches with live progress bar; "
                 "auto-retry 2× per failed batch — failures logged, never silent."),
                ("✏️", "Editable UI Prompts",
                 "Users customise generation rules in the Streamlit UI at run-time — "
                 "no code changes required for new constraints."),
                ("📥", "Instant Download",
                 "Results exported as clean CSV or validated JSON; "
                 "session state preserves data for preview and multi-format download."),
            ],
            "outcome": "🏆  Outcome: Working end-to-end generator — Cluster CSV & ProductVendor JSON "
                       "delivered at 10,000+ rows/run with field-level validation.",
        },
        # ── PHASE 2 ──────────────────────────────────────────────────────────
        {
            "num":    "Phase 2",
            "title":  "PII-Grounded Improvement",
            "status": "🔄  IN PROGRESS",
            "status_color": C_ORANGE,
            "icon":   "🔗",
            "color":  C_ACCENT,
            "bg":     phase_bg[1],
            "what": "Pull anonymised reference data from live Databricks tables to "
                    "anchor the LLM's output — replacing pure randomness with "
                    "distributions that mirror real-world data.",
            "bullets": [
                ("🗄️", "Databricks Table Integration",
                 "Query production tables (with PII stripped/masked) to extract "
                 "value distributions, code lists, and valid ID ranges used in the real system."),
                ("📊", "Distribution-Aware Prompting",
                 "Inject real statistical profiles (e.g. top store numbers, active dept codes, "
                 "season frequency) into prompts so generated data follows actual business patterns."),
                ("🔐", "Privacy-Safe Reference Extraction",
                 "PII fields (customer names, addresses, identifiers) are hashed or "
                 "aggregated before injection — only schema shapes & value ranges are used."),
                ("🧹", "Data Profiler Integration",
                 "profiler.py analyses existing synthetic & real datasets, computing "
                 "cardinality, null rates, and format patterns to guide prompt refinement."),
                ("🔁", "Feedback Loop",
                 "Generated batches are compared against real distributions; "
                 "divergence flags trigger prompt adjustment in subsequent runs."),
                ("📈", "Quality Uplift",
                 "Synthetic data now reflects realistic store-cluster ratios, "
                 "seasonal biases, and vendor activity rates seen in production."),
            ],
            "outcome": "🎯  Goal: Synthetic records indistinguishable from real data in terms of "
                       "distribution, cardinality, and business-rule compliance.",
        },
        # ── PHASE 3 ──────────────────────────────────────────────────────────
        {
            "num":    "Phase 3",
            "title":  "Pattern & Relationship Mining",
            "status": "🔜  UP NEXT",
            "status_color": C_PURPLE,
            "icon":   "🕸️",
            "color":  C_PURPLE,
            "bg":     phase_bg[2],
            "what": "Move beyond per-field statistics — identify cross-entity "
                    "correlations, temporal patterns, and causal links so the agent "
                    "generates truly production-representative datasets.",
            "bullets": [
                ("🔍", "Cross-Field Correlation Analysis",
                 "Mine co-occurrence patterns between fields: e.g. which DEPT_NO values "
                 "consistently appear with specific CSTD_GRADE codes or PHASE_STARTDATE windows."),
                ("📅", "Temporal Pattern Detection",
                 "Identify seasonality, launch cycles, and date-range clusters in "
                 "PHASE_STARTDATE and season fields to generate time-coherent synthetic series."),
                ("🕸️", "Entity Relationship Mapping",
                 "Model relationships between Products ↔ Vendors ↔ Seasons ↔ Stores so that "
                 "generated records honour referential integrity across data types."),
                ("🤖", "Pattern-Conditioned Prompting",
                 "Encode discovered patterns as structured constraints in LLM prompts — "
                 "moving from 'valid format' to 'realistic joint distribution'."),
                ("📐", "Schema Inference from DDL",
                 "Automatically derive generation rules from Databricks table schemas "
                 "and data contracts — reducing manual prompt engineering effort."),
                ("🚀", "Production-Ready Datasets",
                 "Output datasets usable directly in ML training, BI prototyping, "
                 "and integration testing without additional transformation or cleaning."),
            ],
            "outcome": "🌟  Vision: Fully autonomous synthetic data that mirrors production "
                       "complexity — including multi-entity joins, seasonal trends, and rare-event rates.",
        },
    ]

    for i, phase in enumerate(phases):
        tile_left = Inches(0.22) + i * (TILE_W + GAP)

        # Card background
        add_rect(slide, tile_left, PHASE_TOP, TILE_W, PHASE_H,
                 fill_color=phase["bg"],
                 line_color=phase["color"], line_width=Pt(1.5))

        # Top colour band
        add_rect(slide, tile_left, PHASE_TOP, TILE_W, Inches(0.04),
                 fill_color=phase["color"])

        # Left accent bar
        add_rect(slide, tile_left, PHASE_TOP, Inches(0.04), PHASE_H,
                 fill_color=phase["color"])

        # Phase number + icon
        add_text_box(slide, f"{phase['icon']}  {phase['num']}",
                     tile_left + Inches(0.12), PHASE_TOP + Inches(0.07),
                     TILE_W - Inches(0.2), Inches(0.32),
                     font_size=Pt(10), bold=True, color=phase["color"])

        # Status badge
        badge_w = Inches(1.38)
        badge_h = Inches(0.26)
        badge_l = tile_left + TILE_W - badge_w - Inches(0.12)
        badge_t = PHASE_TOP + Inches(0.09)
        add_rect(slide, badge_l, badge_t, badge_w, badge_h,
                 fill_color=RGBColor(0x0A, 0x0E, 0x1A),
                 line_color=phase["status_color"], line_width=Pt(1))
        add_text_box(slide, phase["status"],
                     badge_l, badge_t, badge_w, badge_h,
                     font_size=Pt(7.8), bold=True, color=phase["status_color"],
                     align=PP_ALIGN.CENTER)

        # Title
        add_text_box(slide, phase["title"],
                     tile_left + Inches(0.12), PHASE_TOP + Inches(0.38),
                     TILE_W - Inches(0.2), Inches(0.38),
                     font_size=Pt(13), bold=True, color=C_WHITE)

        # Separator line
        add_rect(slide, tile_left + Inches(0.12), PHASE_TOP + Inches(0.76),
                 TILE_W - Inches(0.25), Inches(0.02),
                 fill_color=phase["color"])

        # "What" description
        add_text_box(slide, phase["what"],
                     tile_left + Inches(0.12), PHASE_TOP + Inches(0.80),
                     TILE_W - Inches(0.2), Inches(0.72),
                     font_size=Pt(9.2), color=C_LIGHT_TEXT, italic=True)

        # Bullet points
        for j, (icon_b, title_b, desc_b) in enumerate(phase["bullets"]):
            by = PHASE_TOP + Inches(1.54) + j * Inches(0.63)

            # Small accent pill for bullet title
            add_rect(slide, tile_left + Inches(0.12), by,
                     TILE_W - Inches(0.25), Inches(0.23),
                     fill_color=RGBColor(0x0E, 0x14, 0x28))
            add_text_box(slide, f"{icon_b}  {title_b}",
                         tile_left + Inches(0.15), by,
                         TILE_W - Inches(0.32), Inches(0.23),
                         font_size=Pt(8.8), bold=True, color=phase["color"])

            # Description
            add_text_box(slide, desc_b,
                         tile_left + Inches(0.15), by + Inches(0.23),
                         TILE_W - Inches(0.32), Inches(0.38),
                         font_size=Pt(8.2), color=C_LIGHT_TEXT)

        # Outcome bar
        outcome_t = PHASE_TOP + PHASE_H - Inches(0.40)
        add_rect(slide, tile_left + Inches(0.04), outcome_t,
                 TILE_W - Inches(0.08), Inches(0.36),
                 fill_color=RGBColor(0x0A, 0x0E, 0x1A),
                 line_color=phase["color"], line_width=Pt(0.8))
        add_text_box(slide, phase["outcome"],
                     tile_left + Inches(0.1), outcome_t + Inches(0.03),
                     TILE_W - Inches(0.22), Inches(0.32),
                     font_size=Pt(8), bold=True, color=phase["color"])

    # ── Bottom bar ─────────────────────────────────────────────────────────────
    add_rect(slide, Inches(0), Inches(7.44), SLIDE_W, Inches(0.06),
             fill_color=C_ACCENT)
    add_rect(slide, Inches(0), Inches(7.26), SLIDE_W, Inches(0.18),
             fill_color=RGBColor(0x0A, 0x14, 0x26))
    add_text_box(slide,
                 "Powered by Databricks · Claude Sonnet · Streamlit  ·  "
                 "Phase 1 ✅ Complete  |  Phase 2 🔄 In Progress  |  Phase 3 🔜 Upcoming",
                 Inches(0), Inches(7.27), SLIDE_W, Inches(0.18),
                 font_size=Pt(8.2), color=C_MUTED, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    prs = new_prs()
    print("Building roadmap slide...")
    build_roadmap_slide(prs)
    print("  ✔  Problem Statement & Evolution slide built")

    output_path = r"C:\Users\P9193015\Downloads\CODE\DATA GEN\Roadmap_Evolution_Slide.pptx"
    prs.save(output_path)
    print(f"\n✅  Saved to:\n   {output_path}")

